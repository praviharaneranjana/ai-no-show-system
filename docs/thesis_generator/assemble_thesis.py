import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.append('docs/thesis_generator')
import format_helpers
import chapter1
import chapter2
import chapter3
import chapter4
import chapter5
import chapter6
import back_matter

print("Assembling complete dissertation document...")

doc = docx.Document()

# Configure Normal style
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(12)
normal_style.font.color.rgb = RGBColor(0, 0, 0)
normal_style.paragraph_format.line_spacing = 1.5
normal_style.paragraph_format.space_after = Pt(6)
normal_style.paragraph_format.space_before = Pt(0)

# Section 1: Cover Page
s_cover = doc.sections[0]
format_helpers.set_page_margins(s_cover)
s_cover.header.is_linked_to_previous = False
s_cover.footer.is_linked_to_previous = False
for p in s_cover.footer.paragraphs:
    p.text = ""

format_helpers.add_p(doc, space_before=72)
format_helpers.add_p(doc, "SHAI: AN EXPLAINABLE MACHINE LEARNING AND DETERMINISTIC ORCHESTRATION PLATFORM FOR OUTPATIENT NO-SHOW PREDICTION, CAPACITY OPTIMIZATION, AND AUTONOMOUS RESCHEDULING IN SRI LANKAN HOSPITALS",
                     align=WD_ALIGN_PARAGRAPH.CENTER, space_before=36, space_after=144, bold=True, size=16)

format_helpers.add_p(doc, "PRAVIHARA NERANJANA ABEGUNASEKARA", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=36, space_after=18, bold=True, size=14)
format_helpers.add_p(doc, "Bachelor of Science (Honours) in Software Engineering", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=144, bold=False, size=14)

format_helpers.add_p(doc, "Department of Software Engineering and Information Systems", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=6, bold=False, size=14)
format_helpers.add_p(doc, "Faculty of Computing", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6, bold=False, size=14)
format_helpers.add_p(doc, "NSBM Green University", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6, bold=False, size=14)
format_helpers.add_p(doc, "Sri Lanka", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=48, bold=False, size=14)

format_helpers.add_p(doc, "October 2026", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=0, bold=False, size=14)

# Section 2: Preliminaries (Roman numerals ii through xi)
s_prelim = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
format_helpers.set_page_margins(s_prelim)
s_prelim.header.is_linked_to_previous = False
s_prelim.footer.is_linked_to_previous = False

sectPr_prelim = s_prelim._sectPr
pgNumType_prelim = OxmlElement('w:pgNumType')
pgNumType_prelim.set(qn('w:fmt'), 'lowerRoman')
pgNumType_prelim.set(qn('w:start'), '1')
sectPr_prelim.append(pgNumType_prelim)

# Page i unprinted
s_prelim.different_first_page_header_footer = True
for p in s_prelim.first_page_footer.paragraphs:
    p.text = ""

f_p_prelim = s_prelim.footer.paragraphs[0]
f_p_prelim.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_fp = f_p_prelim.add_run()
run_fp.font.name = 'Times New Roman'
run_fp.font.size = Pt(12)
run_fp._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w')))
run_fp._r.append(parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w')))
run_fp._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w')))
run_fp._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w')))

# Appendix II: Inner Title Page
format_helpers.add_p(doc, space_before=48)
format_helpers.add_p(doc, "SHAI: AN EXPLAINABLE MACHINE LEARNING AND DETERMINISTIC ORCHESTRATION PLATFORM FOR OUTPATIENT NO-SHOW PREDICTION, CAPACITY OPTIMIZATION, AND AUTONOMOUS RESCHEDULING IN SRI LANKAN HOSPITALS",
                     align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=72, bold=True, size=16)

format_helpers.add_p(doc, "A thesis submitted to NSBM Green University for the degree of\nBachelor of Science (Honours) in Software Engineering",
                     align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=48, bold=False, size=14)

format_helpers.add_p(doc, "By", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=12, bold=False, size=14)
format_helpers.add_p(doc, "PRAVIHARA NERANJANA ABEGUNASEKARA", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=96, bold=True, size=14)

format_helpers.add_p(doc, "Department of Software Engineering and Information Systems", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=6, bold=False, size=14)
format_helpers.add_p(doc, "Faculty of Computing", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6, bold=False, size=14)
format_helpers.add_p(doc, "NSBM Green University", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6, bold=False, size=14)
format_helpers.add_p(doc, "Sri Lanka", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=48, bold=False, size=14)
format_helpers.add_p(doc, "October 2026", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=0, bold=False, size=14)

doc.add_page_break()

# Appendix III: Declaration (Page ii)
format_helpers.add_p(doc, "DECLARATION", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=36, bold=True, size=14)

format_helpers.add_p(doc, "I declare that the content of this undergraduate thesis titled \"SHAI: An Explainable Machine Learning and Deterministic Orchestration Platform for Outpatient No-Show Prediction, Capacity Optimization, and Autonomous Rescheduling in Sri Lankan Hospitals\" is my own work and this dissertation does not incorporate without acknowledgement any material previously submitted for any other degree in any university or institution of higher learning.",
                     space_after=48)

p_sig = doc.add_paragraph()
p_sig.paragraph_format.line_spacing = 1.5
p_sig.paragraph_format.space_before = Pt(36)
p_sig.paragraph_format.space_after = Pt(12)
r_sig1 = p_sig.add_run("Signature: .....................................................")
r_sig1.font.name = 'Times New Roman'
r_sig1.font.size = Pt(12)
r_tab = p_sig.add_run("\t\t\t")
r_date = p_sig.add_run("Date: .....................................")
r_date.font.name = 'Times New Roman'
r_date.font.size = Pt(12)

p_name = doc.add_paragraph()
p_name.paragraph_format.line_spacing = 1.5
p_name.paragraph_format.space_after = Pt(72)
r_name = p_name.add_run("Name: Pravihara Neranjana Abegunasekara")
r_name.font.name = 'Times New Roman'
r_name.font.size = Pt(12)

format_helpers.add_p(doc, "Signature of the Supervisor", space_before=36, space_after=18, bold=True, size=12)
format_helpers.add_p(doc, ".........................................................................", space_after=12)
format_helpers.add_p(doc, "Mr. Anton Jayakody", bold=True, space_after=4)
format_helpers.add_p(doc, "Principal Supervisor", space_after=4)
format_helpers.add_p(doc, "Senior Lecturer", space_after=4)
format_helpers.add_p(doc, "Department of Software Engineering and Information Systems", space_after=4)
format_helpers.add_p(doc, "Faculty of Computing", space_after=4)
format_helpers.add_p(doc, "NSBM Green University, Sri Lanka", space_after=24)

doc.add_page_break()

# Acknowledgement (Page iii)
format_helpers.add_p(doc, "ACKNOWLEDGEMENT", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=24, bold=True, size=14)

format_helpers.add_p(doc, "I would like to express my deepest gratitude to my principal research supervisor, Mr. Anton Jayakody, for his invaluable guidance, technical critique, and intellectual encouragement throughout the development and evaluation of this research. His constructive feedback during system reviews and architectural refinements was vital in shaping the rigorous methodological foundation of this work.")

format_helpers.add_p(doc, "I extend my sincere appreciation to the academic faculty, lecturers, and administrative staff of the Faculty of Computing, NSBM Green University, Sri Lanka, for fostering an inspiring educational environment and providing the computing resources and facilities necessary to complete this degree programme.")

format_helpers.add_p(doc, "I am also profoundly grateful to the clinical, operational, and administrative staff at private healthcare institutions in Colombo, particularly those who provided qualitative domain insights into outpatient consultation operations, multi-channel channelling systems, doctor punctuality dynamics, and front-desk appointment triage workflows. Their real-world perspectives provided the foundational context for modeling operational constraints within the Sri Lankan healthcare landscape.")

format_helpers.add_p(doc, "Finally, my deepest appreciation goes to my family, colleagues, and fellow peers for their unwavering patience, moral encouragement, and steadfast belief in my aspirations throughout the challenging journey of completing this dissertation.")

doc.add_page_break()

# Abstract (Page iv)
format_helpers.add_p(doc, "ABSTRACT", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=24, bold=True, size=14)

format_helpers.add_p(doc, "Outpatient clinic absenteeism presents a severe operational and economic challenge to healthcare institutions globally and within Sri Lanka's dual public-private healthcare system. In Sri Lankan private hospitals, where specialist consultations operate through multi-channel channelling systems under rigid time constraints, unexpected patient no-shows lead to idle clinical faculty, fragmented resource utilization, and lost consultation revenue. While machine learning classification models have been widely proposed in the international literature, existing systems treat prediction as an isolated classification task, failing to address probability miscalibration, operational capacity overbooking, and patient-centric intervention policies. Furthermore, existing research rarely incorporates local operational friction factors such as multi-modal transit distances, Poya holiday disruptions, monsoon weather shocks, and trilingual patient communication.", space_after=12)

format_helpers.add_p(doc, "This dissertation presents SHAI (Smart Hospital AI Operations Platform), an explainable machine learning and deterministic operations platform designed to predict, manage, and mitigate outpatient no-shows. The platform implements an end-to-end architecture comprising an engineered chronological data pipeline, a tuned XGBoost gradient boosting classifier calibrated via isotonic regression, a TreeSHAP explainability engine, an exact Poisson-binomial dynamic programming overbooking optimizer, and an automated ask-first rescheduling state machine operating in English, Sinhala, and Tamil. The system integrates a transparent heuristic action recommendation layer, sensitive clinical specialty privacy masking for mental health departments, and a mathematically rigorous randomized hold-out protocol enabling true intention-to-treat clinical evaluation.", space_after=12)

format_helpers.add_p(doc, "Evaluated on an out-of-sample chronological test set (N = 9,385), the calibrated XGBoost model achieved a receiver operating characteristic area under the curve (ROC-AUC) of 0.672 (95% CI: [0.658, 0.686]), an average precision (PR-AUC) of 0.266, and a calibrated Brier score of 0.1237, outperforming logistic regression, decision trees, and random forests. At the top 20% risk tier, the model achieved a recall of 36.3% and a lift of 1.82x over random outreach. On a real-world Kaggle medical appointment benchmark (N = 110,527), the pipeline demonstrated robust generalizability with an ROC-AUC of 0.736 and a lift of 1.904x. A 90-day operational replay across 9,266 hospital bookings demonstrated that risk-ranked targeting successfully identified 47.4% of all no-shows compared to 27.4% under random targeting, yielding an empirical lift of LKR 228,600 over the evaluation window.", space_after=12)

format_helpers.add_p(doc, "Economic sensitivity modeling based on Sri Lankan clinical baseline literature (Rajasuriya et al., 2010) and empirical telco tariffs (LKR 5.38 blended contact cost) established a platform break-even consultation fee of LKR 3,961, demonstrating that predictive outreach delivers positive operational ROI across prevailing private specialist consultation rates (LKR 4,500+ yielding 1.14x ROI). The complete software system comprises 32 PostgreSQL relational tables, robust role-based access control, and an automated test suite of 233 backend and 31 frontend passing tests, providing a robust, transparent, and clinically deployable decision-support asset for hospital operations.", space_after=12)

doc.add_page_break()

# Appendix IV: Table of Contents (Page v)
format_helpers.add_p(doc, "TABLE OF CONTENTS", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=24, bold=True, size=14)

toc_items = [
    ("Declaration", "ii"),
    ("Acknowledgement", "iii"),
    ("Abstract", "iv"),
    ("Table of Contents", "v"),
    ("List of Figures", "vii"),
    ("List of Tables", "ix"),
    ("List of Abbreviations", "xi"),
    ("1. INTRODUCTION", "1"),
    ("  1.1 Background and Context", "1"),
    ("  1.2 Problem Statement", "4"),
    ("  1.3 Research Justification and Motivation", "6"),
    ("  1.4 Research Questions and Hypotheses", "8"),
    ("  1.5 Scope and Delimitations of the Study", "10"),
    ("  1.6 Outline of the Dissertation", "11"),
    ("2. OBJECTIVES", "13"),
    ("  2.1 General Objective", "13"),
    ("  2.2 Specific Objectives", "13"),
    ("    2.2.1 Predictive Modeling and Statistical Calibration Objectives", "13"),
    ("    2.2.2 Clinical Explainability and Interpretability Objectives", "14"),
    ("    2.2.3 Capacity Optimization and Queue Engineering Objectives", "14"),
    ("    2.2.4 Deterministic Multi-Turn Patient Orchestration Objectives", "15"),
    ("    2.2.5 Clinical Privacy, Fairness, and Scientific Hold-Out Objectives", "15"),
    ("    2.2.6 Economic Sensitivity and Decision-Theoretic ROI Objectives", "16"),
    ("3. LITERATURE REVIEW", "17"),
    ("  3.1 Determinants and Dynamics of Outpatient Absenteeism", "17"),
    ("  3.2 Machine Learning in Healthcare Appointment Scheduling", "20"),
    ("  3.3 The Imperative of Clinical Probability Calibration", "24"),
    ("  3.4 Explainable AI (XAI) in Clinical Decision-Support Systems", "27"),
    ("  3.5 Patient Engagement Interventions and Communication Modalities", "30"),
    ("  3.6 The Sri Lankan Healthcare Baseline: A Critical Analysis of Rajasuriya et al. (2010)", "33"),
    ("  3.7 Operations Research and Dynamic Overbooking Models in Healthcare", "36"),
    ("  3.8 Summary and Identification of Research Gaps", "39"),
    ("4. METHODOLOGY", "42"),
    ("  4.1 Research Design and System Architecture", "42"),
    ("  4.2 Synthetic Data Engineering and Healthcare Data Modeling", "46"),
    ("    4.2.1 Data Generation Assumptions and Sri Lankan Environmental Features", "46"),
    ("    4.2.2 Relational Database Schema and Multi-Tenancy Architecture", "49"),
    ("  4.3 Feature Engineering and Chronological Data Partitioning", "53"),
    ("  4.4 Predictive Machine Learning Pipeline and Bayesian Hyperparameter Tuning", "56"),
    ("  4.5 Probability Calibration and Rank-Based Operational Risk Tiers", "60"),
    ("  4.6 Global and Local Explainability Engine (TreeSHAP)", "63"),
    ("  4.7 Transparent Heuristic Action Recommendation Layer", "65"),
    ("  4.8 Exact Poisson-Binomial Dynamic Programming Overbooking Optimizer", "68"),
    ("  4.9 Ask-First Rescheduling Conversational State Machine", "72"),
    ("  4.10 Multi-Lingual Natural Language Engine and Clinical Privacy Safeguards", "76"),
    ("  4.11 Scientific Evaluation Protocol: Randomized Outreach Hold-Out Design", "79"),
    ("  4.12 Software Quality Assurance, Testing Framework, and Verification Strategy", "82"),
    ("5. RESULTS", "85"),
    ("  5.1 Exploratory Data Analysis and Feature Distribution Verification", "85"),
    ("  5.2 Baseline Model Comparison and Benchmark Performance", "88"),
    ("  5.3 Validation on Real-World Kaggle Medical Appointments Benchmark", "92"),
    ("  5.4 XGBoost Bayesian Tuning and Chronological Generalization", "94"),
    ("  5.5 Probability Calibration Curves and Reliability Verification", "97"),
    ("  5.6 Empirical Hypothesis Testing: Feature Odds Ratios and Significance", "101"),
    ("  5.7 Global and Local TreeSHAP Explainability Results", "104"),
    ("  5.8 Subgroup Fairness and Demographic Parity Audit", "108"),
    ("  5.9 Operational Decision Support and Capacity Management Performance", "112"),
    ("  5.10 End-to-End Clinical Triage and Master Scheduling Implementation", "116"),
    ("  5.11 Multi-Turn Conversational Rescheduling and Staff-in-the-Loop Evaluation", "120"),
    ("  5.12 90-Day Operational Database Replay and Ranking Efficacy Lift", "124"),
    ("  5.13 Decision-Theoretic Economic Evaluation and Sensitivity Analysis", "127"),
    ("  5.14 Software Verification, Reliability Metrics, and Test Execution Results", "131"),
    ("6. DISCUSSION AND CONCLUSIONS", "134"),
    ("  6.1 Synthesis of Findings and Research Contributions", "134"),
    ("  6.2 Comparison with Prior Literature and Empirical Reconciliation", "137"),
    ("  6.3 Practical Implications for Sri Lankan Hospital Management", "140"),
    ("  6.4 Clinical Safety, Privacy Governance, and Ethical Considerations", "143"),
    ("  6.5 Threats to Validity and Study Limitations", "145"),
    ("  6.6 Recommendations for Future Research", "148"),
    ("  6.7 Concluding Remarks", "150"),
    ("REFERENCES", "152"),
    ("APPENDICES", "162"),
    ("  Appendix A: Core Database Schema and Entity-Relationship Definitions", "162"),
    ("  Appendix B: Machine Learning Feature Dictionary and Variable Specifications", "166"),
    ("  Appendix C: Conversational State Machine Dialogue Templates (EN, SI, TA)", "169"),
    ("  Appendix D: Hospital Administration and Role-Based Access Audit Log Samples", "172"),
]

t_toc = doc.add_table(rows=len(toc_items), cols=2)
t_toc.autofit = False
for idx, (item, page) in enumerate(toc_items):
    row = t_toc.rows[idx]
    c0 = row.cells[0]
    c1 = row.cells[1]
    c0.width = Inches(5.3)
    c1.width = Inches(0.9)
    p0 = c0.paragraphs[0]
    p0.paragraph_format.line_spacing = 1.15
    p0.paragraph_format.space_before = Pt(1)
    p0.paragraph_format.space_after = Pt(1)
    r0 = p0.add_run(item)
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11)
    if not item.startswith("  "):
        r0.bold = True
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_before = Pt(1)
    p1.paragraph_format.space_after = Pt(1)
    r1 = p1.add_run(page)
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    if not item.startswith("  "):
        r1.bold = True

doc.add_page_break()

# Appendix V: List of Figures (Page vii)
format_helpers.add_p(doc, "LIST OF FIGURES", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=24, bold=True, size=14)

figures_list = [
    ("Figure 4.1", "High-Level Architecture of the SHAI Healthcare Operations Platform", "44"),
    ("Figure 4.2", "Entity-Relationship Data Model of the Hospital Scheduling Schema", "51"),
    ("Figure 4.3", "Finite-State Transition Architecture for Autonomous Conversational Rescheduling", "74"),
    ("Figure 5.1", "Exploratory Distribution of Absenteeism Rates Across Environmental and Behavioral Dimensions", "86"),
    ("Figure 5.2", "Receiver Operating Characteristic (ROC) Comparison Curves on Real-World Kaggle Benchmark", "93"),
    ("Figure 5.3", "Cumulative Gain and Lift Curves for the Tuned XGBoost Model", "96"),
    ("Figure 5.4", "Out-of-Sample Reliability Diagram for Raw and Isotonically Calibrated XGBoost Probabilities", "99"),
    ("Figure 5.5", "Global TreeSHAP Summary Beeswarm Plot of Predictive Feature Attributions", "105"),
    ("Figure 5.6", "Mean Absolute SHAP Feature Importance Hierarchy for the Tuned Classifier", "107"),
    ("Figure 5.7", "Multi-Dimensional Subgroup Fairness Evaluation Across Demographic and Logistical Strata", "110"),
    ("Figure 5.8", "SHAI Capacity Overview Dashboard Displaying Punctuality Tracking and Live Operational Metrics", "113"),
    ("Figure 5.9", "15-Minute Slot Overbooking Risk Heatmap Matrix for Real-Time Clinic Sessions", "114"),
    ("Figure 5.10", "What-If Studio Dynamic Overbooking Optimization Curve and Session Value Analysis", "115"),
    ("Figure 5.11", "Multi-Horizon Capacity Forecast with 80% Poisson-Binomial Confidence Bounds", "116"),
    ("Figure 5.12", "Appointments Master Schedule Interface with Calibrated Risk Triage and Recommended Actions", "117"),
    ("Figure 5.13", "Patient Intake Wizard Step 1: Patient Identity, Comorbidity Registration, and Risk Preview", "118"),
    ("Figure 5.14", "Patient Intake Wizard Step 2: Commute Distance Calculation and Geospatial Transit Mapping", "119"),
    ("Figure 5.15", "Patient Intake Wizard Step 3: Clinical Specialty Selection and Dynamic Re-Scoring", "120"),
    ("Figure 5.16", "Patient Intake Wizard Step 4: Multi-Channel Booking Confirmation and Review", "121"),
    ("Figure 5.17", "Prediction Hub Daily Batch Scoring and Outreach Drawer Execution Interface", "122"),
    ("Figure 5.18", "Autonomous Outreach Hub WhatsApp Thread with Ask-First Flow and Staff Approval Gate", "123"),
    ("Figure 5.19", "Post-Missed Appointment Recall Queue Benchmarked Against Historical Clinical Efficacy", "124"),
    ("Figure 5.20", "90-Day Operational Historical Replay Lift Comparing Risk-Ranked Outreach to Random Contact", "126"),
    ("Figure 5.21", "Interactive Outreach ROI Simulation and Decision-Theoretic Financial Evaluation Studio", "129"),
    ("Figure 5.22", "Staff and Access Governance Console Enforcing Multi-Tenant Role-Based Access Control", "133"),
]

t_lof = doc.add_table(rows=len(figures_list), cols=2)
t_lof.autofit = False
for idx, (f_num, f_title, f_page) in enumerate(figures_list):
    row = t_lof.rows[idx]
    c0 = row.cells[0]
    c1 = row.cells[1]
    c0.width = Inches(5.3)
    c1.width = Inches(0.9)
    p0 = c0.paragraphs[0]
    p0.paragraph_format.line_spacing = 1.15
    p0.paragraph_format.space_before = Pt(1)
    p0.paragraph_format.space_after = Pt(1)
    r0 = p0.add_run(f"{f_num}\t{f_title}")
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_before = Pt(1)
    p1.paragraph_format.space_after = Pt(1)
    r1 = p1.add_run(f_page)
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)

doc.add_page_break()

# Appendix VI: List of Tables (Page ix)
format_helpers.add_p(doc, "LIST OF TABLES", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=24, bold=True, size=14)

tables_list = [
    ("Table 3.1", "Empirical Baseline Synthesis from Rajasuriya et al. (2010) Outpatient Reminder Study", "34"),
    ("Table 4.1", "Relational Database Schema Table Catalog and Functional Specifications", "50"),
    ("Table 4.2", "Engineered Predictive Feature Schema and Variable Descriptions", "54"),
    ("Table 4.3", "Rule-Based Heuristic Action Recommendation Decision Logic", "67"),
    ("Table 4.4", "Software Architecture Component Catalog and Technology Stack Versions", "83"),
    ("Table 5.1", "Empirical Classification Performance Across Four Candidate Baseline Models on Synthetic Data", "89"),
    ("Table 5.2", "Benchmark Model Performance Evaluation on Real-World Kaggle Medical Appointments Dataset", "92"),
    ("Table 5.3", "Optuna Bayesian Optimization Search Space and Optimal Hyperparameters for Tuned XGBoost", "95"),
    ("Table 5.4", "Probability Calibration Brier Scores and Risk Tier Distribution Metrics", "100"),
    ("Table 5.5", "Multivariable Logistic Regression Odds Ratios, 95% Confidence Intervals, and p-values", "102"),
    ("Table 5.6", "Mean Absolute TreeSHAP Attribution Scores for Top Predictive Features", "106"),
    ("Table 5.7", "Subgroup Fairness Audit Across Demographic, Financial, and Logistical Categories", "109"),
    ("Table 5.8", "Historical 90-Day Operational Replay Performance: Risk-Ranked Policy vs Random Outreach", "125"),
    ("Table 5.9", "Decision-Theoretic Sensitivity Matrix: Annual Net Economic Benefit (LKR) Across Efficacy and Contact Ratios", "128"),
    ("Table 5.10", "Economic Assumption Parameter Register and Evidence Sourcing Summary", "130"),
    ("Table 5.11", "Automated Software Quality Assurance Test Suite Execution Summary", "132"),
]

t_lot = doc.add_table(rows=len(tables_list), cols=2)
t_lot.autofit = False
for idx, (t_num, t_title, t_page) in enumerate(tables_list):
    row = t_lot.rows[idx]
    c0 = row.cells[0]
    c1 = row.cells[1]
    c0.width = Inches(5.3)
    c1.width = Inches(0.9)
    p0 = c0.paragraphs[0]
    p0.paragraph_format.line_spacing = 1.15
    p0.paragraph_format.space_before = Pt(1)
    p0.paragraph_format.space_after = Pt(1)
    r0 = p0.add_run(f"{t_num}\t{t_title}")
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_before = Pt(1)
    p1.paragraph_format.space_after = Pt(1)
    r1 = p1.add_run(t_page)
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)

doc.add_page_break()

# Appendix VII: List of Abbreviations (Page xi)
format_helpers.add_p(doc, "LIST OF ABBREVIATIONS", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=24, bold=True, size=14)

abbreviations = [
    ("API", "Application Programming Interface"),
    ("AUC", "Area Under the Curve"),
    ("CI", "Confidence Interval"),
    ("DP", "Dynamic Programming"),
    ("EPF", "Employees' Provident Fund"),
    ("ETF", "Employees' Trust Fund"),
    ("GBDT", "Gradient Boosted Decision Trees"),
    ("GIS", "Geographic Information System"),
    ("HIS", "Hospital Information System"),
    ("HTTP", "Hypertext Transfer Protocol"),
    ("IEEE", "Institute of Electrical and Electronics Engineers"),
    ("ITT", "Intention-To-Treat"),
    ("JSON", "JavaScript Object Notation"),
    ("KPI", "Key Performance Indicator"),
    ("LKR", "Sri Lankan Rupee"),
    ("LLM", "Large Language Model"),
    ("LR", "Logistic Regression"),
    ("ML", "Machine Learning"),
    ("NIC", "National Identity Card"),
    ("NHSL", "National Hospital of Sri Lanka"),
    ("NLP", "Natural Language Processing"),
    ("OPD", "Outpatient Department"),
    ("OR", "Odds Ratio"),
    ("OSRM", "Open Source Routing Machine"),
    ("PR-AUC", "Precision-Recall Area Under the Curve"),
    ("RBAC", "Role-Based Access Control"),
    ("RCT", "Randomized Controlled Trial"),
    ("REST", "Representational State Transfer"),
    ("RF", "Random Forest"),
    ("ROC", "Receiver Operating Characteristic"),
    ("ROI", "Return on Investment"),
    ("SD", "Standard Deviation"),
    ("SHAP", "Shapley Additive Explanations"),
    ("SMS", "Short Message Service"),
    ("SQL", "Structured Query Language"),
    ("TreeSHAP", "Tree-based Shapley Additive Explanations"),
    ("UHID", "Universal Hospital Identification"),
    ("UI", "User Interface"),
    ("UUID", "Universally Unique Identifier"),
    ("VIVA", "Oral Examination (Viva Voce)"),
    ("XAI", "Explainable Artificial Intelligence"),
    ("XGBoost", "Extreme Gradient Boosting"),
]

t_abbr = doc.add_table(rows=len(abbreviations), cols=2)
t_abbr.autofit = False
for idx, (abbr, desc) in enumerate(abbreviations):
    row = t_abbr.rows[idx]
    c0 = row.cells[0]
    c1 = row.cells[1]
    c0.width = Inches(1.8)
    c1.width = Inches(4.4)
    p0 = c0.paragraphs[0]
    p0.paragraph_format.line_spacing = 1.15
    p0.paragraph_format.space_before = Pt(1)
    p0.paragraph_format.space_after = Pt(1)
    r0 = p0.add_run(abbr)
    r0.bold = True
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_before = Pt(1)
    p1.paragraph_format.space_after = Pt(1)
    r1 = p1.add_run(desc)
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)

# SECTION 3: MAIN BODY CHAPTERS (Arabic numerals starting at 1 at 1 INTRODUCTION)
s_body = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
format_helpers.set_page_margins(s_body)
s_body.header.is_linked_to_previous = False
s_body.footer.is_linked_to_previous = False

sectPr_body = s_body._sectPr
pgNumType_body = OxmlElement('w:pgNumType')
pgNumType_body.set(qn('w:fmt'), 'decimal')
pgNumType_body.set(qn('w:start'), '1')
sectPr_body.append(pgNumType_body)

f_p_body = s_body.footer.paragraphs[0]
f_p_body.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_fb = f_p_body.add_run()
run_fb.font.name = 'Times New Roman'
run_fb.font.size = Pt(12)
run_fb._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w')))
run_fb._r.append(parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w')))
run_fb._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w')))
run_fb._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w')))

print("Writing Chapter 1...")
chapter1.write_chapter1(doc,
                        lambda t: format_helpers.add_h1(doc, t),
                        lambda t: format_helpers.add_h2(doc, t),
                        lambda t: format_helpers.add_h3(doc, t),
                        lambda t, **kwargs: format_helpers.add_p(doc, t, **kwargs))
doc.add_page_break()

print("Writing Chapter 2...")
chapter2.write_chapter2(doc,
                        lambda t: format_helpers.add_h1(doc, t),
                        lambda t: format_helpers.add_h2(doc, t),
                        lambda t: format_helpers.add_h3(doc, t),
                        lambda t, **kwargs: format_helpers.add_p(doc, t, **kwargs))
doc.add_page_break()

print("Writing Chapter 3...")
chapter3.write_chapter3(doc,
                        lambda t: format_helpers.add_h1(doc, t),
                        lambda t: format_helpers.add_h2(doc, t),
                        lambda t: format_helpers.add_h3(doc, t),
                        lambda t, **kwargs: format_helpers.add_p(doc, t, **kwargs),
                        lambda n, t: format_helpers.add_caption_table(doc, n, t),
                        format_helpers.format_custom_table)
doc.add_page_break()

print("Writing Chapter 4...")
chapter4.write_chapter4(doc,
                        lambda t: format_helpers.add_h1(doc, t),
                        lambda t: format_helpers.add_h2(doc, t),
                        lambda t: format_helpers.add_h3(doc, t),
                        lambda t: format_helpers.add_h4(doc, t),
                        lambda t, **kwargs: format_helpers.add_p(doc, t, **kwargs),
                        lambda n, t: format_helpers.add_caption_figure(doc, n, t),
                        lambda n, t: format_helpers.add_caption_table(doc, n, t),
                        lambda p, **kwargs: format_helpers.add_figure_image(doc, p, **kwargs),
                        format_helpers.format_custom_table)
doc.add_page_break()

print("Writing Chapter 5...")
chapter5.write_chapter5(doc,
                        lambda t: format_helpers.add_h1(doc, t),
                        lambda t: format_helpers.add_h2(doc, t),
                        lambda t: format_helpers.add_h3(doc, t),
                        lambda t, **kwargs: format_helpers.add_p(doc, t, **kwargs),
                        lambda n, t: format_helpers.add_caption_figure(doc, n, t),
                        lambda n, t: format_helpers.add_caption_table(doc, n, t),
                        lambda p, **kwargs: format_helpers.add_figure_image(doc, p, **kwargs),
                        format_helpers.format_custom_table)
doc.add_page_break()

print("Writing Chapter 6...")
chapter6.write_chapter6(doc,
                        lambda t: format_helpers.add_h1(doc, t),
                        lambda t: format_helpers.add_h2(doc, t),
                        lambda t: format_helpers.add_h3(doc, t),
                        lambda t, **kwargs: format_helpers.add_p(doc, t, **kwargs))
doc.add_page_break()

print("Writing References and Appendices...")
back_matter.write_back_matter(doc,
                              lambda t: format_helpers.add_h1(doc, t),
                              lambda t: format_helpers.add_h2(doc, t),
                              lambda t: format_helpers.add_h3(doc, t),
                              lambda t, **kwargs: format_helpers.add_p(doc, t, **kwargs),
                              lambda n, t: format_helpers.add_caption_table(doc, n, t),
                              format_helpers.format_custom_table)

output_path = "SHAI_Final_Thesis_Pravihara_Abegunasekara.docx"
print(f"Saving final dissertation document to {output_path}...")
doc.save(output_path)
print("Dissertation generation completed successfully!")
